"""Create the explanatory Word dossier from executed project outputs."""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
D=json.loads((ROOT/'site/data.json').read_text())
SC=pd.read_csv(ROOT/'outputs/scenarios.csv')
O=D['overview']; S=D['optimization']
doc=Document();sec=doc.sections[0]
sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(1.85);sec.left_margin=Cm(2.15);sec.right_margin=Cm(2.15)
styles=doc.styles
normal=styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(10);normal.font.color.rgb=RGBColor(36,44,54)
normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.17
for name,size,before,after in [('Title',23,0,11),('Heading 1',15,15,7),('Heading 2',11.7,12,5)]:
    st=styles[name];st.font.name='Aptos Display';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.font.bold=name!='Title';st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after)
    st.paragraph_format.keep_with_next=True
# Word's built-in Title style may carry a theme-colored bottom border.
tp=styles['Title']._element.get_or_add_pPr()
for old in tp.findall(qn('w:pBdr')):tp.remove(old)
none=OxmlElement('w:pBdr');edge=OxmlElement('w:bottom');edge.set(qn('w:val'),'nil');none.append(edge);tp.append(none)

def P(text='',style=None):return doc.add_paragraph(text,style)
def H(text,level=1):doc.add_heading(text,level)
def boldline(lead,rest):
    p=doc.add_paragraph();r=p.add_run(lead);r.bold=True;p.add_run(rest)
def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),fill);tcPr.append(shd)
def borders(cell):
    tcPr=cell._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right'):
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bd.append(e)
    tcPr.append(bd)
def pad(cell):
    tc=cell._tc;tcPr=tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
    for edge in ('top','left','bottom','right'):
        e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
    tcPr.append(mar)
def tbl(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Cm(w)
    for i,v in enumerate(headers):
        t.rows[0].cells[i].text=str(v);shade(t.rows[0].cells[i],'20354D')
    for j,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val)
            if j%2:shade(cells[i],'F1F5F8')
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            borders(c);pad(c);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.08
                if ci and len(headers)>2:p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
                for run in p.runs:
                    run.font.size=Pt(8.6);run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(27,39,50)
                    if ri==0:run.bold=True
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
def fmt(x,d=0):return f'{x:,.{d}f}'.replace(',','X').replace('.',',').replace('X','.')

doc.add_paragraph('Dossier del Proyecto 7 de Levante Ferries','Title')
P('Sostenibilidad y descarbonización marítima para decisiones de operaciones y finanzas')
P('Álvaro Martínez Lacasta  |  Portfolio Dirección + Datos + IA  |  Septiembre de 2026')
boldline('Naturaleza del caso. ','Levante Ferries es una naviera ficticia. Los datos, costes, buques y resultados proceden de un modelo reproducible; las reglas europeas citadas son reales. No se presenta como trabajo realizado para ninguna naviera real.')
H('Resumen ejecutivo')
P(f'El proyecto une viaje, ruta, buque, combustible, demanda, costes y emisiones para elegir medidas de descarbonización que conserven capacidad y nivel de servicio. Entre enero y agosto de 2026 el caso simulado contiene {fmt(O["voyages"])} viajes, {fmt(O["passengers"])} pasajeros y {fmt(O["ttw_t"])} t de CO₂e emitidas a bordo. El margen modelizado tras energía y derechos EUA asciende a {fmt(O["margin_eur"]/1e6,2)} M€.')
P(f'El optimizador propone una alternativa con {fmt(S["optimized_ttw_t"])} t CO₂e, una reducción de {fmt(100*(1-S["optimized_ttw_t"]/S["baseline_ttw_t"]),1)} %, y un coste {fmt(S["savings_eur"]/1e6,2)} M€ menor. Esta diferencia es un resultado condicional del modelo. La recomendación ejecutiva es validar horarios, conexión OPS, combustible certificado y costes de implantación antes de contratar cambios.')
tbl(['Indicador del caso','Valor','Lectura'],[
 ('CO₂e TTW',f'{fmt(O["ttw_t"])} t','Inventario físico con CH₄ y N₂O'),
 ('CO₂e WTW',f'{fmt(O["wtw_t"])} t','Ciclo de vida, red OPS incluida'),
 ('Coste EU ETS',f'{fmt(O["ets_eur"]/1e6,2)} M€','EUA a 75 €/t, precio supuesto'),
 ('Intensidad FuelEU proxy',f'{fmt(O["fueleu_g_mj"],2)} g/MJ','Límite modelizado 89,34 g/MJ')],[4.6,3.0,8.5])

H('1 Pregunta de dirección y arquitectura')
P('La pregunta es cómo reducir emisiones mientras la naviera mantiene ingresos, capacidad y salidas. Cada viaje registra ruta y distancia, buque y su compatibilidad, energía marítima y en puerto, demanda de pasajeros y carga, ingresos, coste fijo y alcance territorial. La función de cálculo produce emisiones, energía, coste EU ETS y margen. Los cambios de velocidad, mezcla de combustible, OPS, demanda o precio EUA actualizan la misma cadena de cálculo.')
P('Los proyectos anteriores aportan conceptos de operación, predicción, tarifa y control financiero. Los datos de este proyecto son nuevos y sintéticos: las claves de viaje de los repositorios previos no se unen sin correspondencia validada. El agente corporativo del Proyecto 6 podría explicar estas métricas y citar supuestos a partir de vistas agregadas, sujeto a supervisión humana.')

H('2 Datos y control de calidad')
P('Se generan 4.860 piernas de viaje en ocho rutas y diez buques ficticios, con fechas entre el 1 de enero y el 31 de agosto de 2026. Hay rutas entre puertos de la UE y dos corredores con tercer país. Todas las capacidades, demandas, tarifas, consumos y factores energéticos son hipótesis. La semilla 20260923 permite reproducir el mismo conjunto.')
P('La base valida identificadores únicos, ausencia de valores nulos, demanda no superior a capacidad y consistencia entre buque y ruta. Los agregados dividen sumas homogéneas; no promedian intensidades de viajes. La ruta y el buque forman una pareja de decisión del optimizador; éste mantiene cada salida y su demanda, sin cambiar frecuencias.')

H('3 Contabilidad de carbono y KPIs')
P('El inventario tank to wake se calcula con CO₂ + 28 × CH₄ + 265 × N₂O, en kg CO₂e, mediante factores supuestos. Well to wake usa la energía por combustible y sus factores de ciclo de vida; para el inventario de OPS se asigna un factor de red hipotético de 55 g/MJ. Son lecturas alternativas de una actividad; no se suman entre sí.')
tbl(['KPI','Fórmula y unidad'],[
 ('CO₂e TTW','Σ [CO₂ kg + 28 × CH₄ kg + 265 × N₂O kg]; kg o t'),
 ('CO₂e WTW','Σ MJ de cada combustible × g CO₂e/MJ ÷ 1.000 + OPS de red; kg'),
 ('Intensidad por milla','Σ CO₂e TTW kg ÷ Σ millas náuticas; kg/nm'),
 ('Intensidad por pasajero','Σ CO₂e TTW kg ÷ Σ pasajeros; kg/pasajero'),
 ('Intensidad por carga','Σ CO₂e TTW kg ÷ Σ metros lineales; kg/m lineal'),
 ('Consumo y energía','MJ de combustible y kg por MJ/PCI; OPS MJ aparte'),
 ('Margen','Σ ingresos − Σ (coste fijo + energía + EUA); €')],[4.6,11.5])
P('Las intensidades por pasajero y carga describen un servicio compartido, no asignan el total de emisiones a cada producto por separado. La ocupación afecta al denominador por pasajero y a los ingresos; en este supuesto no modifica directamente la energía requerida para navegar.')

H('4 EU ETS y FuelEU Maritime')
P('EU ETS aplica al caso hipotético de buques ≥5.000 GT. La cobertura geográfica simplificada es 100 % en trayectos entre puertos de la UE/EEE y 50 % con un tercer país. En 2026 se incluyen CO₂, CH₄ y N₂O. El coste devengado es toneladas CO₂e TTW cubiertas × precio EUA supuesto. Las emisiones de 2026 requieren derechos entregados en 2027. [E1]')
P('FuelEU mide intensidad de energía well to wake por buque y año. El caso emplea 91,16 × (1 − 0,02) = 89,3368 g CO₂e/MJ para 2025–2029. La energía eléctrica suministrada en puerto entra en el denominador de la proxy con factor regulatorio cero en el numerador, mientras el inventario global sí refleja el factor de red supuesto. La obligación OPS para pasajeros depende de puerto y condiciones desde 2030. [E2, E3]')
P('El gap se calcula como intensidad modelizada menos límite. El balance mostrado en el dashboard es un proxy energético y no una sanción. Faltan factores certificados, ciclo anual completo, trazabilidad del biofuel, excepciones aplicables, pooling, banking, borrowing y verificación. No se afirma cumplimiento legal. Los buques de MGO del caso base superan el umbral proxy; los de GNL permanecen por debajo bajo los factores supuestos.')

H('5 Escenarios de transición')
P('Cada escenario recalcula los mismos viajes. La velocidad del 94 % supone energía por milla proporcional a velocidad al cuadrado. El biofuel sustituye parcialmente al combustible principal y está condicionado a certificado; OPS solo se aplica a rutas marcadas como disponibles. La simulación de demanda no altera consumos ni emisiones si la oferta de salidas y buques permanece fija.')
tbl(['Escenario','CO₂e TTW t','Cambio','Margen M€'],[
 (r['scenario'],fmt(r['ttw_t']),fmt(r['delta_ttw_pct'],1)+' %',fmt(r['margin_eur']/1e6,2)) for _,r in SC.iterrows()],[6.9,3.2,2.7,3.3])
P('La cartera combinada logra una reducción mayor y un margen ligeramente inferior al base con esos precios de energía. El escenario OPS mejora la intensidad FuelEU proxy mediante el factor regulatorio cero, aunque el coste de electricidad supuesto puede superar el combustible sustituido. Estas sensibilidades no incluyen inversión en conexión, reformas del buque ni cambios comerciales inducidos por horarios.')

H('6 Motor de optimización')
P('Se definen 540 candidatos en diez parejas ruta–buque, variando velocidad, combustible compatible, proporción bio y OPS disponible. Una variable binaria elige un candidato por pareja. La función objetivo minimiza coste fijo operativo + energía + derechos EUA. La demanda y los ingresos quedan fijos por viaje, las frecuencias y capacidades se conservan y el tiempo de mar no supera el máximo supuesto.')
P(f'Las restricciones adicionales limitan emisiones TTW a 88 % del caso base e intensidad FuelEU proxy por buque a 89,3368 g/MJ. El solver devuelve estado óptimo en esta ejecución. La alternativa elegida emite {fmt(S["optimized_ttw_t"])} t frente a {fmt(S["baseline_ttw_t"])} t, y el coste cae de {fmt(S["baseline_cost_eur"]/1e6,2)} a {fmt(S["optimized_cost_eur"]/1e6,2)} M€. La combinación óptima no es necesariamente la que más descarboniza: minimiza el coste dentro de los límites fijados.')
P('El modelo no asigna buques nuevos, no cancela salidas y no incluye CAPEX, subvenciones, contratos, incertidumbre de precios ni restricciones de atraque minuto a minuto. Con datos reales debería resolverse de nuevo tras incorporar forecast, política tarifaria y coste financiero.')

H('7 Qué contienen los siete notebooks')
tbl(['Cuaderno','Resultado principal'],[
 ('01 Data Foundation','Datos sintéticos, granos, dimensiones y validaciones.'),
 ('02 Carbon Accounting','TTW, WTW, gases, energía y denominadores.'),
 ('03 Eco Efficiency Analysis','Comparativa de rutas y buques con ocupación y margen.'),
 ('04 EU ETS & FuelEU Maritime','Derechos EUA, intensidad por buque y gaps proxy.'),
 ('05 Decarbonization Scenarios','Sensibilidades de velocidad, biofuel, OPS y demanda.'),
 ('06 Optimization Engine','Elección binaria, restricciones y conciliación.'),
 ('07 Executive Sustainability Model','Matriz margen–emisiones y prioridades directivas.')],[5.9,10.2])

H('8 Decisiones y hoja de ruta')
P('Primero, usar el dashboard para detectar rutas o buques con bajo margen y alta intensidad, mirando también frecuencia y demanda. Segundo, probar velocidad y OPS únicamente en horarios y puertos viables. Tercero, negociar suministro bio con trazabilidad y estudiar coste total incluyendo inversiones. Cuarto, integrar salidas reales, forecast, precios y finanzas en claves de servicio y periodos conciliados. Finalmente, recalcular un año completo con factores y verificadores antes de cualquier declaración regulatoria.')
P('El site dark incluye Executive Overview, Fleet, Routes, Energy Transition, EU ETS & FuelEU y un simulador interactivo. El repositorio queda estructurado para GitHub y la página preparada para incorporarse más adelante al dominio del portfolio; no se ha publicado en este trabajo.')

H('Fuentes públicas y procedencia')
P('[E1] Comisión Europea, Reducing emissions from the shipping sector. climate.ec.europa.eu; enlace completo en el README.')
P('[E2] Agencia Europea de Seguridad Marítima, FuelEU Maritime Regulation. emsa.europa.eu; enlace completo en el README.')
P('[E3] Comisión Europea, Questions and answers on Regulation 2023/1805. transport.ec.europa.eu; enlace completo en el archivo de fuentes.')
P('[S1] Baleària, comunicaciones sobre bioGNL y electrificación: referencia sectorial pública. No se importan buques, rutas, costes ni emisiones reales a la base ficticia.')

footer=sec.footer.paragraphs[0];footer.text='Levante Ferries  |  Caso sintético  |  Álvaro Martínez Lacasta';footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
for run in footer.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor(100,115,130)
out=ROOT/'Dossier_Proyecto_7_Sustainability_Decarbonization_Levante_Ferries.docx'
doc.save(out);print(out)
